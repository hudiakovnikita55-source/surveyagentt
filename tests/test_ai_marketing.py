import csv
import json
from types import SimpleNamespace

import pytest

from surveyagent.answering.heuristic import HeuristicAnswerer
from surveyagent.answering.llm import build_schema, render_questionnaire
from surveyagent.answering.runner import generate_responses
from surveyagent.merge import _translate, merge, write_merged
from surveyagent.personas import generate_personas, load_personas, preferred_language
from surveyagent.personas.geo import normalize_country
from surveyagent.questionnaire import END, Question, Questionnaire, load_questionnaire, normalize_answer

LANGS = ("en", "pl", "ru")
BOOST = {"PL": 4, "LV": 5, "EE": 5}


@pytest.fixture(scope="module")
def forms():
    return {lang: load_questionnaire(f"surveys/ai_marketing/{lang}.yaml") for lang in LANGS}


@pytest.fixture(scope="module")
def marketers():
    return generate_personas(100, seed=37, audience="marketing", country_boost=BOOST)


# --------------------------------------------------------------------------- questionnaire logic

def test_language_versions_are_aligned(forms):
    en = forms["en"]
    for lang, q in forms.items():
        assert q.language == lang.upper()
        assert [x.id for x in q.questions] == [f"P{i}" for i in range(1, 16)]
        for a, b in zip(q.questions, en.questions):
            assert (a.type, len(a.options), len(a.rows), a.has_other, a.required, a.topic) == \
                   (b.type, len(b.options), len(b.rows), b.has_other, b.required, b.topic)
            assert [(a.options.index(k), v) for k, v in a.go_to.items()] == \
                   [(b.options.index(k), v) for k, v in b.go_to.items()]
            assert [a.options.index(k) for k in a.exclusive] == [b.options.index(k) for k in b.exclusive]


def test_path_follows_go_to(forms):
    en = forms["en"]
    ids = lambda answers: [q.id for q in en.path(answers)]
    assert ids({"P1": "No"}) == ["P1"]
    assert ids({"P1": "Yes", "P2": "No"}) == ["P1", "P2"]
    assert ids({"P1": "Yes", "P2": "Yes", "P7": "No"}) == ["P1", "P2", "P3", "P4", "P5", "P6", "P7", "P12", "P13",
                                                           "P14", "P15"]
    assert len(ids({"P1": "Yes", "P2": "Yes", "P7": "Yes"})) == 15
    assert en.conditional_ids() == {f"P{i}" for i in range(2, 16)}
    assert "logic: \"No\" -> skip to [P12]" in render_questionnaire(en)


def test_logic_is_validated():
    with pytest.raises(ValueError):  # unknown option
        Question(id="a", title="A", type="single_choice", options=["Yes", "No"], go_to={"Maybe": END})
    with pytest.raises(ValueError):  # go_to only on single choice
        Question(id="a", title="A", type="checkboxes", options=["x", "y"], go_to={"x": END})
    with pytest.raises(ValueError):  # backwards jump
        Questionnaire(title="t", questions=[Question(id="a", title="A", type="short_text"),
                                            Question(id="b", title="B", type="single_choice", options=["x", "y"],
                                                     go_to={"x": "a"})])


def test_exclusive_option(forms):
    q = forms["en"].question("P13")
    only = q.exclusive[0]
    assert normalize_answer(q, [only]) == [only]
    with pytest.raises(ValueError):
        normalize_answer(q, [only, q.options[0]])


# --------------------------------------------------------------------------- marketing personas

def test_default_personas_are_unchanged():
    committed = json.loads(open("data/personas.json", encoding="utf-8").read())
    assert [p.to_dict() for p in generate_personas(100, seed=42)] == committed


def test_marketing_audience(marketers):
    again = generate_personas(100, seed=37, audience="marketing", country_boost=BOOST)
    assert [p.to_dict() for p in again] == [p.to_dict() for p in marketers]
    active = [p for p in marketers if p.marketing["active"]]
    assert len(active) >= 85
    non_users = [p for p in active if not p.marketing["uses_ai"]]
    assert 5 <= len(non_users) <= 20
    for p in non_users:
        assert not p.uses_ai and p.ai["tools"] == [] and p.marketing["risk_practices"] == ["not_applicable"]
    for p in active:
        assert p.marketing["setting"] in ("agency", "in_house", "freelance", "own_business", "other")
        assert set(p.marketing["tasks"]) == {"audience_research", "strategy", "segmentation", "copywriting", "visuals",
                                             "ads", "community", "reporting"}
    assert sum(p.country == "Poland" for p in marketers) >= 8


def test_preferred_language(marketers):
    for p in marketers:
        lang = preferred_language(p, ["EN", "PL", "RU"])
        if "Polish" in p.native_languages:
            assert lang == "PL"
        elif "Russian" in p.native_languages:
            assert lang == "RU"
        elif p.country != "Poland":
            assert lang == "EN"
    assert preferred_language(marketers[0], ["PL"]) in ("PL",)


# --------------------------------------------------------------------------- offline answers

def test_offline_answers_follow_logic_and_profile(forms, marketers):
    en = forms["en"]
    answerer = HeuristicAnswerer(en)
    for p in marketers:
        answers = answerer.answer(p)
        shown = {q.id for q in en.path(answers)}
        for q in en.questions:
            if q.id in shown:
                normalize_answer(q, answers[q.id])
            else:
                assert answers[q.id] is None
        m = p.marketing
        assert answers["P1"] == ("Yes" if m["consents"] else "No")
        if not m["consents"]:
            continue
        assert answers["P2"] == ("Yes" if m["active"] else "No")
        if not m["active"]:
            continue
        assert normalize_country(answers["P3"]) == p.country
        assert answers["P7"] == ("Yes" if m["uses_ai"] else "No")
        only = en.question("P13").exclusive[0]
        assert (answers["P13"] == [only]) == (not m["uses_ai"])
        if m["uses_ai"]:
            for row, task in zip(en.question("P10").rows, m["tasks"]):
                assert (answers["P10"][row] == "I do not perform this type of task") == (m["tasks"][task] == "no")
        if m["setting"] == "other":
            assert answers["P4"] == "Other marketing-related role (please specify)"


def test_offline_answers_do_not_depend_on_language(forms, marketers):
    """Same persona, same seed: the PL/RU answers are the EN answers in another language (text aside)."""
    def other_as_marker(value):  # "Other" free text is written in the questionnaire's language
        if isinstance(value, dict):
            return "OTHER"
        if isinstance(value, list):
            return [v if isinstance(v, str) else "OTHER" for v in value]
        return value

    en = forms["en"]
    expected = {p.id: HeuristicAnswerer(en).answer(p) for p in marketers}
    for lang in ("pl", "ru"):
        form = forms[lang]
        for p in marketers:
            answers = HeuristicAnswerer(form).answer(p)
            for sq, rq in zip(form.questions, en.questions):
                if sq.type not in ("short_text", "paragraph"):
                    assert other_as_marker(_translate(answers[sq.id], sq, rq)) == \
                           other_as_marker(expected[p.id][rq.id]), (lang, p.id, sq.id)


# --------------------------------------------------------------------------- LLM mode with logic

def test_schema_allows_null_only_for_skippable_questions(forms):
    schema = build_schema(forms["pl"])["properties"]
    assert "anyOf" not in schema["P1"]
    assert schema["P10"]["anyOf"][1] == {"type": "null"}
    assert schema["P10"]["anyOf"][0]["type"] == "object"
    assert schema["skipped"]["items"]["enum"] == ["P15"]


class LogicClient:
    """A model that says it does not use AI but still fills in P8-P11 (which must then be dropped)."""

    def __init__(self, form, personas):
        self.form, self.by_name = form, {p.full_name: p for p in personas}
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self._create))
        self.messages = self.beta.messages

    def _create(self, **params):
        name = params["messages"][0]["content"].rsplit("as ", 1)[1].rstrip(".")
        answers = HeuristicAnswerer(self.form).answer(self.by_name[name])
        data = {"skipped": []}
        for q in self.form.questions:
            value = answers[q.id]
            if q.type == "grid" and value is not None:
                value = {f"r{i}": value[row] for i, row in enumerate(q.rows, 1)}
            elif isinstance(value, list):
                value = [v for v in value if isinstance(v, str)]
            elif isinstance(value, dict):
                value = "__other__"
            data[q.id] = value
            if q.has_other:
                data[f"{q.id}__other"] = "coś innego"
        data["P7"] = "Nie"
        data["P8"] = "Codziennie"
        return SimpleNamespace(stop_reason="end_turn", model=params["model"],
                               content=[SimpleNamespace(type="text", text=json.dumps(data))],
                               usage=SimpleNamespace(input_tokens=1, output_tokens=1, cache_read_input_tokens=0,
                                                     cache_creation_input_tokens=0))


def test_llm_answers_are_cut_to_the_path(tmp_path, forms, marketers):
    form = forms["pl"]
    people = [p for p in marketers if p.marketing["consents"] and p.marketing["active"]][:3]
    records = generate_responses(form, people, tmp_path / "r.json", mode="llm", client=LogicClient(form, marketers),
                                 log=lambda *_: None)
    for r in records:
        a = r["answers"]
        assert a["P7"] == "Nie"
        assert all(a[f"P{i}"] is None for i in range(8, 12))
        assert a["P13"] is not None and r["synthetic"] is True


# --------------------------------------------------------------------------- merging

def test_merge_json_and_csv_exports_agree(tmp_path, forms, marketers):
    inputs_json, inputs_csv = [], []
    for lang in LANGS:
        form = forms[lang]
        group = [p for p in marketers if preferred_language(p, ["EN", "PL", "RU"]) == lang.upper()]
        out = tmp_path / f"s_{lang}.json"
        generate_responses(form, group, out, mode="offline", log=lambda *_: None)
        inputs_json.append((form, out))
        inputs_csv.append((form, out.with_suffix(".csv")))
    ref, from_json = merge(inputs_json)
    _, from_csv = merge(inputs_csv)
    assert ref.language == "EN" and len(from_json) == 100
    key = lambda r: r["respondent_id"]
    assert sorted(from_json, key=key) == sorted(from_csv, key=key)
    assert {r["language"] for r in from_json} == {"EN", "PL", "RU"}
    assert all(r["synthetic"] == "yes" for r in from_json)
    statuses = {r["status"] for r in from_json}
    assert statuses <= {"complete", "ended_at_P1", "ended_at_P2"} and "complete" in statuses
    pl_row = next(r for r in from_json if r["language"] == "PL" and r["status"] == "complete")
    assert pl_row["P1"] == "Yes" and pl_row["P3_country"]


def test_merge_reads_a_real_google_export(tmp_path, forms):
    """Headers with "P7." prefixes, checkbox answers joined with ", " even inside options, an "only this"
    conflict, free-text countries."""
    pl = forms["pl"]
    header = ["Sygnatura czasowa"] + [f"{q.id}. {q.title}" for q in pl.questions if q.type != "grid"]
    grid_cols = [f"{q.title} [{row}]" for q in pl.questions if q.type == "grid" for row in q.rows]
    p9, p13 = pl.question("P9"), pl.question("P13")
    values = {"P1": "Tak", "P2": "Tak", "P3": "polska ", "P4": "Inna rola związana z marketingiem (jaka?)",
              "P5": "4–6 lat", "P6": "50–249 osób", "P7": "Tak", "P8": "Codziennie",
              "P9": f"{p9.options[0]}, {p9.options[1]}, {p9.options[5]}",
              "P13": f"{p13.options[0]}, {p13.options[7]}", "P14": "Tak", "P15": "Szybciej piszę posty."}
    row = ["2026/10/02 10:15:00 AM GMT+2"] + [values.get(q.id, "") for q in pl.questions if q.type != "grid"]
    row += [q.options[3] for q in pl.questions if q.type == "grid" for _ in q.rows]
    path = tmp_path / "pl_export.csv"
    with open(path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows([header + grid_cols, row])
    warnings = []
    ref, rows = merge([(pl, path)], reference=forms["en"], warnings=warnings)
    assert warnings == []
    r = rows[0]
    assert r["respondent_id"] == "PL-0001" and r["synthetic"] == "no" and r["submitted_at"].startswith("2026/10/02")
    assert r["P3"] == "polska " and r["P3_country"] == "Poland"
    assert r["P4"] == "Other marketing-related role (please specify)"
    assert (r["P9_1"], r["P9_2"], r["P9_3"], r["P9_6"]) == (1, 1, 0, 1)
    assert r["P13_1"] == 1 and r["P13_8"] == 1 and r["P13_conflict"] == 1
    assert r["P10_1"] == forms["en"].question("P10").options[3] and r["P14"] == "Yes"
    codebook = write_merged(tmp_path / "m.csv", ref, rows)
    assert "P13_conflict" in codebook.read_text(encoding="utf-8-sig")


def test_google_sheet_layout_round_trip(tmp_path, forms, marketers):
    """Synthetic answers -> the layout of the form's responses sheet -> CSV export -> read back: same data."""
    import openpyxl
    from surveyagent.merge import collect, write_google_sheet

    versions = [forms["pl"], forms["en"], forms["ru"]]
    inputs = []
    for form in versions:
        group = [p for p in marketers if preferred_language(p, ["EN", "PL", "RU"]) == form.language]
        out = tmp_path / f"s_{form.language}.json"
        generate_responses(form, group, out, mode="offline", log=lambda *_: None)
        inputs.append((form, out))
    ref, rows = collect(inputs)
    path = write_google_sheet(tmp_path / "sheet.xlsx", ref, rows, versions, "Choose language / Wybierz język")
    wb = openpyxl.load_workbook(path)
    fr, cd = wb["Form Responses 1"], wb["Combined data"]
    header = [c.value for c in fr[1]]
    assert header[:3] == ["Timestamp", "Choose language / Wybierz język", "P1. " + forms["pl"].question("P1").title]
    assert len(header) == 2 + 3 * 32 + 1 and header[-1] == "synthetic"
    assert "P10. For which marketing tasks do you use AI? [Audience and market analysis]" in header
    assert [c.value for c in cd[1]] == ["Timestamp", "Language"] + [f"P{i}" for i in range(1, 16)] + ["synthetic"]
    assert {fr.cell(r, 2).value for r in range(2, fr.max_row + 1)} == {"Polski", "English", "Русский"}
    combined = list(cd.iter_rows(min_row=2, values_only=True))
    assert all(r[-1] == "yes" for r in combined)
    grid = next(r[11] for r in combined if r[11])
    assert json.loads(grid) and set(json.loads(grid)) <= set(forms["en"].question("P10").rows)
    assert any(" | " in (r[10] or "") for r in combined)

    export = tmp_path / "export.csv"
    with open(export, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        for row in fr.iter_rows(values_only=True):
            writer.writerow([v.isoformat(sep=" ") if hasattr(v, "isoformat") else ("" if v is None else v)
                             for v in row])
    _, again = collect([(form, export) for form in versions])
    assert sorted((r.language, json.dumps(a, sort_keys=True, ensure_ascii=False)) for r, _, a in again) == \
           sorted((r.language, json.dumps(a, sort_keys=True, ensure_ascii=False)) for r, _, a in rows)


def test_country_normalisation():
    assert normalize_country("Polska") == "Poland"
    assert normalize_country(" великобритания ") == "United Kingdom"
    assert normalize_country("the Netherlands") == "Netherlands"
    assert normalize_country("Rosja") == "Russia"
    assert normalize_country("") is None and normalize_country("Atlantis") is None


def test_committed_marketing_personas_match_generator():
    committed = load_personas("data/marketing_personas.json")
    assert [p.to_dict() for p in committed] == \
           [p.to_dict() for p in generate_personas(100, seed=37, audience="marketing", country_boost=BOOST)]

