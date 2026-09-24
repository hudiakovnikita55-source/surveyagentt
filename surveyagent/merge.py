"""Merge the language versions of one questionnaire into a single analysis table.

Inputs, one per language version, each with that version's questionnaire file:
  - a responses .json written by `answer` (synthetic responses), or
  - a .csv export of real responses (Google Forms -> Responses -> Download CSV, or the Google Sheet as CSV),
    or the responses .csv written by `answer`.
All versions must have the same question ids and the same order of options and grid rows.

Output: one row per respondent, values translated into the labels of a reference version (EN by default):
  respondent_id, language, submitted_at, synthetic, status, then per question:
  single choice -> <id> (+ <id>_other), checkboxes -> <id>_1..<id>_n as 1/0 (+ <id>_other, <id>_conflict when an
  "only this answer" option was ticked together with others), grids -> <id>_1..<id>_n, text -> <id>
  (+ <id>_country with the country name in English for the work-country question).
A codebook CSV next to it lists every column with the question, row and option labels.

`write_google_sheet` writes the same responses in the layout of a Google Sheet linked to ONE form whose first
question picks the language and whose sections hold the language versions: sheet "Form Responses 1"
(Timestamp, language, then every version's columns side by side, filled only for the chosen language) and
sheet "Combined data" (Timestamp, Language, one column per question in the reference labels; multi-select
joined with " | ", grids as JSON). A CSV export of such a sheet can be read back with one --input per version.
"""

from __future__ import annotations

import csv
import difflib
import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from .personas.geo import normalize_country
from .questionnaire import CHOICE_TYPES, END, GRID_TYPES, SCALE_TYPES, Questionnaire, _match_option

TIMESTAMP_HEADERS = {"timestamp", "sygnatura czasowa", "znacznik czasu", "отметка времени", "zeitstempel",
                     "horodateur", "marca temporal", "generated_at"}
# How a version is named in a "choose your language" question.
LANGUAGE_LABELS = {"PL": "Polski", "EN": "English", "RU": "Русский", "UK": "Українська", "DE": "Deutsch",
                   "FR": "Français", "ES": "Español", "IT": "Italiano"}


@dataclass
class Row:
    respondent_id: str
    language: str
    submitted_at: str
    synthetic: bool
    answers: dict  # canonical answers in the labels of the row's own language version


# --------------------------------------------------------------------------- #
# Loading
# --------------------------------------------------------------------------- #

def load_rows(questionnaire: Questionnaire, path: str | Path, warnings: list[str] | None = None) -> list[Row]:
    path = Path(path)
    warnings = [] if warnings is None else warnings
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        ids = [q["id"] for q in data["questionnaire"]["questions"]]
        if ids != [q.id for q in questionnaire.questions]:
            raise ValueError(f"{path}: its questions do not match {questionnaire.source or questionnaire.title}")
        lang = questionnaire.language or data["questionnaire"].get("language") or path.stem
        return [Row(r["persona_id"], lang, r.get("generated_at", ""), r.get("synthetic", True), r["answers"])
                for r in data["responses"]]
    return rows_from_export(questionnaire, path, warnings)


def _norm(text: str) -> str:
    text = re.sub(r"\s+", " ", str(text)).strip().rstrip("*").strip().lower()
    return re.sub(r"^p\d+[a-z]?\s*[.):-]\s*", "", text)  # "P7. Title" -> "title"


def _find(header_norm: list[str], wanted: str) -> int | None:
    key = _norm(wanted)
    if key in header_norm:
        return header_norm.index(key)
    close = difflib.get_close_matches(key, header_norm, n=1, cutoff=0.9)
    return header_norm.index(close[0]) if close else None


def rows_from_export(questionnaire: Questionnaire, path: str | Path, warnings: list[str]) -> list[Row]:
    """Read a CSV export (one column per question, grid rows as "Title [Row]") into canonical answers."""
    with open(path, newline="", encoding="utf-8-sig") as f:
        table = list(csv.reader(f))
    if not table:
        return []
    header, body = table[0], [r for r in table[1:] if any(cell.strip() for cell in r)]
    header_norm = [_norm(h) for h in header]
    lang = questionnaire.language or Path(path).stem

    columns: dict[str, Any] = {}
    for q in questionnaire.questions:
        if q.type in GRID_TYPES:
            cols = {}
            for row in q.rows:
                idx = _find(header_norm, f"{q.title} [{row}]")
                if idx is not None:
                    cols[row] = idx
            if len(cols) < len(q.rows):
                warnings.append(f"{Path(path).name}: {q.id} has {len(cols)} of {len(q.rows)} grid columns")
            columns[q.id] = cols
        else:
            idx = _find(header_norm, q.title)
            if idx is None:
                warnings.append(f"{Path(path).name}: no column for {q.id} ({q.title[:50]})")
            columns[q.id] = idx
    used = {i for v in columns.values() for i in (v.values() if isinstance(v, dict) else [v]) if i is not None}
    special = {h: i for i, h in enumerate(header_norm) if i not in used}
    ts_col = next((i for h, i in special.items() if h in TIMESTAMP_HEADERS), None)
    if ts_col is None and 0 in special.values():
        ts_col = 0
    id_col = special.get("persona_id")
    synthetic_col = special.get("synthetic")

    rows = []
    for n, cells in enumerate(body, 1):
        cells = cells + [""] * (len(header) - len(cells))
        if not any(cells[i].strip() for i in used):
            continue  # a respondent of another language version of the same form
        answers = {}
        for q in questionnaire.questions:
            spec = columns[q.id]
            if q.type in GRID_TYPES:
                value = {}
                for row, idx in spec.items():
                    cell = cells[idx].strip()
                    if cell:
                        parts = _split_multi(cell, q.options)[0] if q.type == "checkbox_grid" else None
                        value[row] = parts if parts is not None else (_match_option(cell, q.options) or cell)
                answers[q.id] = value or None
            else:
                answers[q.id] = parse_cell(q, cells[spec]) if spec is not None else None
        rid = cells[id_col] if id_col is not None else f"{lang}-{n:04d}"
        synthetic = cells[synthetic_col].strip().lower() in ("yes", "true", "1") if synthetic_col is not None else False
        rows.append(Row(rid, lang, cells[ts_col] if ts_col is not None else "", synthetic, answers))
    return rows


def _split_multi(cell: str, options: list[str]) -> tuple[list[str], str]:
    """Checkbox answers are joined with ", " although options may contain commas: match known options first."""
    rest, found = cell, []
    for option in sorted(options, key=len, reverse=True):
        if option in rest:
            found.append(option)
            rest = rest.replace(option, "\x00", 1)
    leftover = ", ".join(p.strip() for p in rest.split("\x00") if p.strip(" ,"))
    return sorted(found, key=options.index), leftover.strip(" ,")


def parse_cell(q, cell: str):
    raw, cell = cell or "", (cell or "").strip()
    if not cell:
        return None
    if q.type in CHOICE_TYPES:
        match = _match_option(cell, q.options)
        return match if match is not None else {"other": cell}
    if q.type == "checkboxes":
        found, leftover = _split_multi(cell, q.options)
        return found + ([{"other": leftover}] if leftover else [])
    if q.type in SCALE_TYPES:
        try:
            return int(float(cell))
        except ValueError:
            return cell
    return raw  # free text exactly as typed (Google keeps trailing spaces too)


# --------------------------------------------------------------------------- #
# Translating between language versions
# --------------------------------------------------------------------------- #

def check_aligned(reference: Questionnaire, other: Questionnaire) -> None:
    problems = []
    if [q.id for q in reference.questions] != [q.id for q in other.questions]:
        problems.append("question ids differ")
    else:
        for a, b in zip(reference.questions, other.questions):
            if (a.type, len(a.options), len(a.rows), a.has_other) != (b.type, len(b.options), len(b.rows), b.has_other):
                problems.append(f"{a.id}: type/options/rows differ")
    if problems:
        raise ValueError(f"{other.language or other.title} does not match {reference.language or reference.title}: "
                         + "; ".join(problems))


def _translate(value: Any, src, ref) -> Any:
    def one(v):
        if isinstance(v, str) and v in src.options:
            return ref.options[src.options.index(v)]
        return v
    if value is None or src is ref:
        return value
    if src.type in GRID_TYPES and isinstance(value, dict):
        return {ref.rows[src.rows.index(r)] if r in src.rows else r:
                ([one(x) for x in c] if isinstance(c, list) else one(c)) for r, c in value.items()}
    if isinstance(value, list):
        return [one(v) for v in value]
    return one(value)


def translate_row(row: Row, source: Questionnaire, reference: Questionnaire) -> dict:
    return {rq.id: _translate(row.answers.get(rq.id), sq, rq) for sq, rq in zip(source.questions, reference.questions)}


# --------------------------------------------------------------------------- #
# Writing
# --------------------------------------------------------------------------- #

def _status(reference: Questionnaire, answers: dict) -> str:
    path = reference.path(answers)
    last = path[-1]
    if last is not reference.questions[-1] and last.jump_for(answers.get(last.id)) == END:
        return f"ended_at_{last.id}"
    return "complete"


def columns_for(reference: Questionnaire) -> list[tuple[str, str, str]]:
    """(column, question id, what it holds) for the merged table and the codebook."""
    cols = []
    for q in reference.questions:
        if q.type == "checkboxes":
            cols += [(f"{q.id}_{i}", q.id, option) for i, option in enumerate(q.options, 1)]
            if q.has_other:
                cols.append((f"{q.id}_other", q.id, "Other (free text)"))
            if q.exclusive:
                cols.append((f"{q.id}_conflict", q.id, f"1 = '{q.exclusive[0]}' ticked together with other options"))
        elif q.type in GRID_TYPES:
            cols += [(f"{q.id}_{i}", q.id, row) for i, row in enumerate(q.rows, 1)]
        else:
            cols.append((q.id, q.id, ""))
            if q.has_other:
                cols.append((f"{q.id}_other", q.id, "Other (free text)"))
            if q.topic == "work_country":
                cols.append((f"{q.id}_country", q.id, "country in English (cleaned)"))
    return cols


def merged_record(reference: Questionnaire, answers: dict) -> dict:
    out = {}
    for q in reference.questions:
        value = answers.get(q.id)
        if q.type == "checkboxes":
            chosen = value or []
            for i, option in enumerate(q.options, 1):
                out[f"{q.id}_{i}"] = "" if value is None else int(option in chosen)
            if q.has_other:
                out[f"{q.id}_other"] = next((v["other"] for v in chosen if isinstance(v, dict)), "")
            if q.exclusive:
                out[f"{q.id}_conflict"] = "" if value is None else int(
                    any(o in chosen for o in q.exclusive) and len(chosen) > 1)
        elif q.type in GRID_TYPES:
            for i, row in enumerate(q.rows, 1):
                cell = (value or {}).get(row, "")
                out[f"{q.id}_{i}"] = "; ".join(cell) if isinstance(cell, list) else cell
        else:
            is_other = isinstance(value, dict)
            out[q.id] = "Other" if is_other else ("" if value is None else value)
            if q.has_other:
                out[f"{q.id}_other"] = value["other"] if is_other else ""
            if q.topic == "work_country":
                out[f"{q.id}_country"] = (normalize_country(value) or "") if isinstance(value, str) else ""
    return out


def collect(inputs: list[tuple[Questionnaire, str | Path]], reference: Questionnaire | None = None,
            warnings: list[str] | None = None) -> tuple[Questionnaire, list[tuple[Row, Questionnaire, dict]]]:
    """Load every input; returns (reference, [(row, its questionnaire, answers in reference labels)])."""
    warnings = [] if warnings is None else warnings
    if reference is None:
        reference = next((q for q, _ in inputs if q.language == "EN"), inputs[0][0])
    out = []
    for questionnaire, path in inputs:
        check_aligned(reference, questionnaire)
        for row in load_rows(questionnaire, path, warnings):
            out.append((row, questionnaire, translate_row(row, questionnaire, reference)))
    return reference, out


def merge(inputs: list[tuple[Questionnaire, str | Path]], reference: Questionnaire | None = None,
          warnings: list[str] | None = None) -> tuple[Questionnaire, list[dict]]:
    """Returns (reference questionnaire, merged rows)."""
    reference, rows = collect(inputs, reference, warnings)
    records = [{"respondent_id": row.respondent_id, "language": row.language, "submitted_at": row.submitted_at,
                "synthetic": "yes" if row.synthetic else "no", "status": _status(reference, answers),
                **merged_record(reference, answers)} for row, _, answers in rows]
    return reference, records


def write_merged(path: str | Path, reference: Questionnaire, records: list[dict]) -> Path:
    """Writes the merged CSV and its codebook (<name>_codebook.csv); returns the codebook path."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = columns_for(reference)
    header = ["respondent_id", "language", "submitted_at", "synthetic", "status"] + [c for c, _, _ in cols]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=header)
        writer.writeheader()
        writer.writerows(records)
    codebook = path.with_name(path.stem + "_codebook.csv")
    with open(codebook, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["column", "question_id", "question", "row_or_option", "values"])
        for name, qid in [("respondent_id", ""), ("language", ""), ("submitted_at", ""), ("synthetic", ""),
                          ("status", "")]:
            values = {"synthetic": "yes = generated by surveyagent, no = real response",
                      "status": "complete | ended_at_<id> (the form's logic ended the survey there)"}.get(name, "")
            writer.writerow([name, qid, "", "", values])
        questions = {q.id: q for q in reference.questions}
        for name, qid, what in cols:
            q = questions[qid]
            if name.endswith(("_other", "_conflict", "_country")):
                values = "text" if not name.endswith("_conflict") else "1 / 0"
            elif q.type == "checkboxes":
                values = "1 = ticked, 0 = not ticked, empty = question not shown"
            elif q.type in GRID_TYPES or q.type in CHOICE_TYPES:
                values = " | ".join(q.options) + (" | Other" if q.has_other and q.type in CHOICE_TYPES else "")
            else:
                values = "text" if q.type not in SCALE_TYPES else f"{q.scale_min}..{q.scale_max}"
            writer.writerow([name, qid, q.title, what, values])
    return codebook


# --------------------------------------------------------------------------- #
# Google-Sheet layout of a single multi-language form
# --------------------------------------------------------------------------- #

def _to_datetime(text: str):
    text = str(text or "").strip()
    for parse in (lambda t: datetime.fromisoformat(t).replace(tzinfo=None),
                  lambda t: datetime.strptime(t.split(" GMT")[0], "%m/%d/%Y %H:%M:%S"),
                  lambda t: datetime.strptime(t.split(" GMT")[0], "%Y/%m/%d %I:%M:%S %p"),
                  lambda t: datetime.strptime(t, "%d.%m.%Y %H:%M:%S")):
        try:
            return parse(text)
        except ValueError:
            continue
    return text or None


def _own_value(q, value):
    """A cell as Google Forms writes it into the responses sheet."""
    if value is None:
        return None
    if isinstance(value, dict) and set(value) == {"other"}:
        return value["other"]
    if isinstance(value, list):
        return ", ".join(v["other"] if isinstance(v, dict) else str(v) for v in value) or None
    return value


def _combined_value(q, value):
    if value is None:
        return None
    if q.type in GRID_TYPES:
        cells = {row: value[row] for row in q.rows if row in value}
        return json.dumps(cells, ensure_ascii=False, separators=(",", ":")) if cells else None
    if isinstance(value, list):
        return " | ".join(v["other"] if isinstance(v, dict) else str(v) for v in value) or None
    if isinstance(value, dict) and set(value) == {"other"}:
        return value["other"]
    return value


def write_google_sheet(path: str | Path, reference: Questionnaire, rows: list[tuple[Row, Questionnaire, dict]],
                       versions: list[Questionnaire], language_question: str = "Language") -> Path:
    """Writes an .xlsx with the sheets "Form Responses 1" and "Combined data" (see the module docstring).

    `versions` gives the order of the language blocks, as in the form. A "synthetic" column is added at the end
    of both sheets (and an explanatory sheet) when any row was generated rather than answered by a person.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    synthetic = any(row.synthetic for row, _, _ in rows)
    ordered = sorted(rows, key=lambda r: (not isinstance(_to_datetime(r[0].submitted_at), datetime),
                                          str(_to_datetime(r[0].submitted_at))))
    wb = Workbook()
    ws = wb.active
    ws.title = "Form Responses 1"
    header = ["Timestamp", language_question]
    for version in versions:
        for q in version.questions:
            if q.type in GRID_TYPES:
                header += [f"{q.id}. {q.title} [{row}]" for row in q.rows]
            else:
                header.append(f"{q.id}. {q.title}")
    ws.append(header + (["synthetic"] if synthetic else []))
    for row, source, _ in ordered:
        line = [_to_datetime(row.submitted_at), LANGUAGE_LABELS.get(row.language, row.language)]
        for version in versions:
            for q in version.questions:
                value = row.answers.get(q.id) if version is source else None
                if q.type in GRID_TYPES:
                    line += [(value or {}).get(r) for r in q.rows]
                else:
                    line.append(_own_value(q, value))
        ws.append(line + (["yes" if row.synthetic else "no"] if synthetic else []))

    wc = wb.create_sheet("Combined data")
    wc.append(["Timestamp", "Language"] + [q.id for q in reference.questions] + (["synthetic"] if synthetic else []))
    for row, _, answers in ordered:
        line = [_to_datetime(row.submitted_at), row.language]
        line += [_combined_value(q, answers.get(q.id)) for q in reference.questions]
        wc.append(line + (["yes" if row.synthetic else "no"] if synthetic else []))

    for sheet in (ws, wc):
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for cell in sheet["A"][1:]:
            cell.number_format = "m/d/yyyy h:mm:ss"
        sheet.column_dimensions["A"].width = 19
    if synthetic:
        note = wb.create_sheet("About this file")
        note.append(["SYNTHETIC DATA / ДАННЫЕ СГЕНЕРИРОВАНЫ"])
        note.append(["Rows with synthetic = yes were generated by surveyagent for fictional personas. "
                     "They are test data for building the analysis, not survey responses."])
        note.append(["Строки с synthetic = yes сгенерированы программой от имени вымышленных персон. "
                     "Это тестовые данные для подготовки анализа, а не ответы участников опроса."])
        note["A1"].font = Font(bold=True, color="9C0006")
    wb.save(path)
    return path
